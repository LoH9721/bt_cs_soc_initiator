# -*- coding: utf-8 -*-
"""CR009 手机 PEPS 自动解闭锁 使用说明 -> docx (宋体黑字)"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\myProject\20260817\bt_cs_soc_initiator_slc_failed_backup\docs\CR009_USAGE_GUIDE.md"
OUT = r"D:\myProject\20260817\bt_cs_soc_initiator_slc_failed_backup\docs\CR009_手机PEPS自动解闭锁使用说明.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
SONG = "宋体"


def set_run_font(run, size, bold=False, mono=False):
    run.font.name = "SimSun"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    if mono:
        rfonts.set(qn("w:ascii"), "Consolas")
        rfonts.set(qn("w:hAnsi"), "Consolas")
        rfonts.set(qn("w:eastAsia"), SONG)
    else:
        rfonts.set(qn("w:ascii"), "SimSun")
        rfonts.set(qn("w:hAnsi"), "SimSun")
        rfonts.set(qn("w:eastAsia"), SONG)


def para(doc, text, size=11, bold=False, align=None, before=0, after=4,
         indent=None, mono=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.25
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    run = p.add_run(text)
    set_run_font(run, size, bold, mono)
    return p


def set_cell(cell, text, size=10, bold=False, width=None):
    if width is not None:
        cell.width = Cm(width)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size, bold)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, rows, col_widths, header_bold=True, font_size=10):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            w = col_widths[j] if j < len(col_widths) else 3.0
            set_cell(table.cell(i, j), val, size=font_size,
                     bold=(i == 0 and header_bold), width=w)
    # fix grid columns
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), col_widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    # explicit table total width (DXA)
    total = int(sum(col_widths) * 567)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    return table


def get_widths(ncols):
    # A4 可用宽度 15.4cm (21 - 2.8*2)
    if ncols == 2:
        return [4.2, 11.2]
    if ncols == 3:
        return [4.0, 5.7, 5.7]
    if ncols == 6:
        return [1.0, 2.4, 3.0, 3.6, 3.2, 2.2]
    return [15.4 / ncols] * ncols


def main():
    doc = Document()
    # A4, margins
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    in_code = False
    code_buf = []
    table_buf = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        if line.startswith("```"):
            if in_code:
                for cl in code_buf:
                    para(doc, cl, size=9.5, after=0, indent=0.4, mono=True)
                para(doc, "", size=4, after=2)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            # 过滤 markdown 表格分隔行 (|---|)
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        if table_buf:
            add_table(doc, table_buf, get_widths(len(table_buf[0])))
            para(doc, "", size=4, after=2)
            table_buf = []
            # 不 continue: 当前行仍需按正文/标题处理
        if line.startswith("# "):
            title = line[2:].strip()
            para(doc, title, size=18, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=10)
        elif line.startswith("## "):
            para(doc, line[3:].strip(), size=14, bold=True, before=14, after=6)
        elif line.startswith("### "):
            para(doc, line[4:].strip(), size=12, bold=True, before=10, after=4)
        elif line.startswith("- "):
            para(doc, "\u2022 " + line[2:].strip(), size=11, after=3, indent=0.4)
        elif line.strip() == "":
            pass
        else:
            para(doc, line.strip(), size=11, after=4)
        i += 1
    if table_buf:
        add_table(doc, table_buf, get_widths(len(table_buf[0])))

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
